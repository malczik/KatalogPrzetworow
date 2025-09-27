# -*- coding: utf-8 -*-
from docx import Document
from docx.shared import Inches
from datetime import datetime

class DocxExporter:
    """Klasa do eksportowania przepisów do DOCX"""
    
    def export_recipe(self, recipe, file_path, calculation_info=None):
        """Eksportuje przepis do pliku DOCX"""
        try:
            doc = Document()
            
            # Nagłówek
            title = doc.add_heading('Przepis na przetwory', 0)
            doc.add_paragraph(f"Wygenerowano: {datetime.now().strftime('%Y-%m-%d %H:%M')}")
            doc.add_paragraph()
            
            # Informacje o przeliczeniu
            if calculation_info:
                doc.add_heading('Informacje o przeliczeniu', level=1)
                for line in calculation_info.split('\n'):
                    if line.strip():
                        doc.add_paragraph(line.strip())
                doc.add_paragraph()
            
            # Parametry przepisu
            doc.add_heading('Parametry przepisu', level=1)
            doc.add_paragraph(f"Pojemność słoików: {recipe.original_jar_capacity} ml")
            doc.add_paragraph(f"Liczba słoików: {recipe.original_jar_count}")
            doc.add_paragraph(f"Całkowita objętość: {recipe.original_jar_capacity * recipe.original_jar_count} ml")
            doc.add_paragraph()
            
            # Składniki
            doc.add_heading('Składniki', level=1)
            for ingredient in recipe.ingredients:
                if ingredient.form:
                    line = f"• {ingredient.name}: {ingredient.amount} {ingredient.unit} ({ingredient.form})"
                else:
                    line = f"• {ingredient.name}: {ingredient.amount} {ingredient.unit}"
                doc.add_paragraph(line)
            
            # Stopka
            doc.add_paragraph()
            doc.add_paragraph("Wygenerowano przez Kalkulator Przetworów")
            
            doc.save(file_path)
            return True
            
        except Exception as e:
            raise Exception(f"Błąd eksportu DOCX: {str(e)}")