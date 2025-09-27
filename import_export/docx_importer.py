# -*- coding: utf-8 -*-
from docx import Document
import re
from core.recipe import Recipe, Ingredient

class DocxImporter:
    """Klasa do importowania przepisów z plików DOCX"""
    
    def import_recipe(self, file_path):
        """Importuje przepis z pliku DOCX"""
        try:
            doc = Document(file_path)
            recipe = Recipe()
            ingredients = []
            
            # Przetwarzaj zarówno akapity jak i tabele
            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    self._process_line(paragraph.text.strip(), recipe, ingredients)
            
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if cell.text.strip():
                            self._process_line(cell.text.strip(), recipe, ingredients)
            
            # Jeśli nie znaleziono składników, spróbuj innej metody
            if not ingredients:
                ingredients = self._extract_ingredients_from_text(doc)
            
            recipe.ingredients = ingredients
            
            # Domyślne wartości jeśli nie znaleziono
            if recipe.original_jar_capacity == 0:
                recipe.original_jar_capacity = 500
            if recipe.original_jar_count == 0:
                recipe.original_jar_count = 5
                
            return recipe
            
        except Exception as e:
            raise Exception(f"Błąd importu DOCX: {str(e)}")
    
    def _process_line(self, line, recipe, ingredients):
        """Przetwarza pojedynczą linię tekstu"""
        line_lower = line.lower()
        
        # Szukaj pojemności
        if any(keyword in line_lower for keyword in ['pojemność', 'pojemnosc', 'ml']):
            numbers = re.findall(r'\d+', line)
            if numbers:
                recipe.original_jar_capacity = float(numbers[0])
        
        # Szukaj liczby słoików
        elif any(keyword in line_lower for keyword in ['słoik', 'sloik', 'słoiki', 'sloiki', 'liczba']):
            numbers = re.findall(r'\d+', line)
            if numbers:
                recipe.original_jar_count = float(numbers[0])
        
        # Próbuj parsować jako składnik
        else:
            ingredient = self._parse_ingredient_line(line)
            if ingredient:
                ingredients.append(ingredient)
    
    def _extract_ingredients_from_text(self, doc):
        """Alternatywna metoda ekstrakcji składników"""
        ingredients = []
        full_text = ""
        
        for paragraph in doc.paragraphs:
            full_text += paragraph.text + "\n"
        
        # Szukaj sekcji ze składnikami
        lines = full_text.split('\n')
        in_ingredients_section = False
        
        for line in lines:
            line = line.strip()
            if not line:
                continue
                
            if any(keyword in line.lower() for keyword in ['składniki', 'skladniki', 'ingredients']):
                in_ingredients_section = True
                continue
                
            if in_ingredients_section:
                ingredient = self._parse_ingredient_line(line)
                if ingredient:
                    ingredients.append(ingredient)
        
        return ingredients
    
    def _parse_ingredient_line(self, line):
        """Parsuje pojedynczą linię na składnik"""
        try:
            # Czyszczenie linii
            line = re.sub(r'[^\w\s\d.,]', ' ', line)
            line = re.sub(r'\s+', ' ', line).strip()
            
            parts = line.split()
            if len(parts) < 2:
                return None
            
            # Szukaj liczby w różnych formatach
            amount = None
            unit = None
            name_parts = []
            
            for i, part in enumerate(parts):
                # Sprawdź czy część jest liczbą
                part_clean = part.replace(',', '.')
                if re.match(r'^\d+\.?\d*$', part_clean):
                    try:
                        amount = float(part_clean)
                        # Spróbuj znaleźć jednostkę w następnych częściach
                        if i + 1 < len(parts):
                            next_part = parts[i + 1].lower()
                            unit = self._normalize_unit(next_part)
                        break
                    except ValueError:
                        continue
                else:
                    name_parts.append(part)
            
            if amount is None:
                # Spróbuj znaleźć liczbę na końcu
                for i in range(len(parts)-1, -1, -1):
                    part_clean = parts[i].replace(',', '.')
                    if re.match(r'^\d+\.?\d*$', part_clean):
                        try:
                            amount = float(part_clean)
                            unit = self._normalize_unit(parts[i-1]) if i > 0 else 'g'
                            name_parts = parts[:i-1] if i > 0 else parts[:i]
                            break
                        except ValueError:
                            continue
            
            if amount is None:
                return None
                
            if unit is None:
                unit = 'g'  # Domyślna jednostka
                
            name = ' '.join(name_parts).strip()
            if not name:
                return None
                
            return Ingredient(name, amount, unit, None)
            
        except Exception:
            return None
    
    def _normalize_unit(self, unit_str):
        """Normalizuje jednostki do standardowych form"""
        unit_str = unit_str.lower()
        
        unit_mapping = {
            'g': ['g', 'gram', 'gramy'],
            'kg': ['kg', 'kilogram', 'kilogramy'],
            'ml': ['ml', 'mililitr', 'mililitry'],
            'l': ['l', 'litr', 'litry'],
            'łyżka': ['łyżka', 'łyżki', 'lyzka', 'lyzki'],
            'łyżeczka': ['łyżeczka', 'łyżeczki', 'lyzeczka', 'lyzeczki'],
            'szklanka': ['szklanka', 'szklanki', 'szkl'],
            'szt': ['szt', 'sztuka', 'sztuki'],
            'ząbek': ['ząbek', 'ząbki'],
            'plaster': ['plaster', 'plastry']
        }
        
        for standard_unit, variants in unit_mapping.items():
            if unit_str in variants:
                return standard_unit
                
        return 'g'  # Domyślna jednostka